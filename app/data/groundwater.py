import csv
import json
import math
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from sqlalchemy import select
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.orm import Session

from app.data.common import sha256_file
from app.models.entities import (
    DataImportRun,
    DataSource,
    GroundwaterObservation,
    Station,
)
from app.schemas.common import MonitoringMode, QualityFlag

AUDIT_VERSION = "gwt-audit-v1"
EXPECTED_HEADERS = ("date", "groundwater table (m)", "station", "peatland")


class GroundwaterDataError(ValueError):
    pass


@dataclass(frozen=True)
class RawGroundwaterRow:
    source_row_number: int
    observed_date: date
    gwt_m: float | None
    station_code: str
    region: str


@dataclass(frozen=True)
class GroundwaterRecord:
    observed_date: date
    gwt_m: float | None
    station_code: str
    region: str
    quality_flag: QualityFlag
    source_row_number: int | None


@dataclass(frozen=True)
class GroundwaterAudit:
    report: dict[str, Any]
    records: tuple[GroundwaterRecord, ...]


def _normalize_headers(values: tuple[Any, ...]) -> tuple[str, ...]:
    return tuple(str(value).strip().lower() for value in values)


def _parse_gwt(value: Any, row_number: int) -> float | None:
    if value == "#N/A" or value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise GroundwaterDataError(f"Unexpected GWT value at Excel row {row_number}: {value!r}")
    if not math.isfinite(value):
        raise GroundwaterDataError(f"Non-finite GWT at Excel row {row_number}")
    return float(value)


def _read_rows(path: Path) -> tuple[list[RawGroundwaterRow], int]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if "gwt" not in workbook.sheetnames:
            raise GroundwaterDataError("Workbook does not contain the required 'gwt' sheet")
        worksheet = workbook["gwt"]
        iterator = worksheet.iter_rows(values_only=True)
        try:
            headers = _normalize_headers(next(iterator))
        except StopIteration as exc:
            raise GroundwaterDataError("The 'gwt' sheet is empty") from exc
        if headers != EXPECTED_HEADERS:
            raise GroundwaterDataError(
                f"Unexpected gwt headers: {headers!r}; expected {EXPECTED_HEADERS!r}"
            )

        rows: list[RawGroundwaterRow] = []
        for row_number, values in enumerate(iterator, start=2):
            raw_date, raw_gwt, raw_station, raw_region = values
            if isinstance(raw_date, (int, float)) and not isinstance(raw_date, bool):
                try:
                    raw_date = from_excel(raw_date, workbook.epoch)
                except (ValueError, OverflowError) as exc:
                    raise GroundwaterDataError(f"Invalid date at Excel row {row_number}") from exc
            if not isinstance(raw_date, (date, datetime)):
                raise GroundwaterDataError(
                    f"Unexpected date at Excel row {row_number}: {raw_date!r}"
                )
            station = str(raw_station).strip().upper() if raw_station is not None else ""
            region = str(raw_region).strip() if raw_region is not None else ""
            if not station or not region:
                raise GroundwaterDataError(f"Missing station or region at Excel row {row_number}")
            observed_date = raw_date.date() if isinstance(raw_date, datetime) else raw_date
            rows.append(
                RawGroundwaterRow(
                    source_row_number=row_number,
                    observed_date=observed_date,
                    gwt_m=_parse_gwt(raw_gwt, row_number),
                    station_code=station,
                    region=region,
                )
            )
        wrc_rows = max(workbook["wrc"].max_row - 1, 0) if "wrc" in workbook.sheetnames else 0
        return rows, wrc_rows
    finally:
        workbook.close()


def audit_groundwater(path: str | Path) -> GroundwaterAudit:
    source_path = Path(path).resolve()
    source_hash = sha256_file(source_path)
    rows, wrc_rows = _read_rows(source_path)
    if sha256_file(source_path) != source_hash:
        raise GroundwaterDataError("Source file changed during audit")
    grouped: dict[tuple[str, date], list[RawGroundwaterRow]] = defaultdict(list)
    station_regions: dict[str, set[str]] = defaultdict(set)
    for row in rows:
        grouped[(row.station_code, row.observed_date)].append(row)
        station_regions[row.station_code].add(row.region)

    conflicting_regions = {
        station: sorted(regions)
        for station, regions in station_regions.items()
        if len(regions) != 1
    }
    if conflicting_regions:
        raise GroundwaterDataError(f"Stations have conflicting regions: {conflicting_regions}")

    records: list[GroundwaterRecord] = []
    duplicate_groups: list[dict[str, Any]] = []
    for (station_code, observed_date), group in sorted(grouped.items()):
        if len(group) > 1:
            duplicate_groups.append(
                {
                    "station_code": station_code,
                    "observed_date": observed_date.isoformat(),
                    "rows": [
                        {"source_row_number": row.source_row_number, "gwt_m": row.gwt_m}
                        for row in group
                    ],
                }
            )
            records.append(
                GroundwaterRecord(
                    observed_date=observed_date,
                    gwt_m=None,
                    station_code=station_code,
                    region=group[0].region,
                    quality_flag=QualityFlag.DUPLICATE,
                    source_row_number=None,
                )
            )
            continue

        row = group[0]
        records.append(
            GroundwaterRecord(
                observed_date=row.observed_date,
                gwt_m=row.gwt_m,
                station_code=row.station_code,
                region=row.region,
                quality_flag=QualityFlag.VALID if row.gwt_m is not None else QualityFlag.MISSING,
                source_row_number=row.source_row_number,
            )
        )

    station_reports: list[dict[str, Any]] = []
    for station_code in sorted(station_regions):
        raw_station_rows = [row for row in rows if row.station_code == station_code]
        station_records = [record for record in records if record.station_code == station_code]
        dates = sorted(record.observed_date for record in station_records)
        valid_values = [
            record.gwt_m
            for record in station_records
            if record.quality_flag is QualityFlag.VALID and record.gwt_m is not None
        ]
        gap_segments = [
            {
                "after": left.isoformat(),
                "before": right.isoformat(),
                "missing_days": (right - left).days - 1,
            }
            for left, right in zip(dates, dates[1:], strict=False)
            if (right - left).days > 1
        ]
        flags = Counter(record.quality_flag.value for record in station_records)
        station_reports.append(
            {
                "station_code": station_code,
                "region": next(iter(station_regions[station_code])),
                "raw_rows": len(raw_station_rows),
                "unique_observations": len(station_records),
                "valid": flags[QualityFlag.VALID.value],
                "missing": flags[QualityFlag.MISSING.value],
                "duplicate_conflicts": flags[QualityFlag.DUPLICATE.value],
                "first_date": dates[0].isoformat(),
                "last_date": dates[-1].isoformat(),
                "min_gwt_m": min(valid_values) if valid_values else None,
                "max_gwt_m": max(valid_values) if valid_values else None,
                "gap_days": sum(segment["missing_days"] for segment in gap_segments),
                "gap_segments": gap_segments,
            }
        )

    flags = Counter(record.quality_flag.value for record in records)
    report: dict[str, Any] = {
        "audit_version": AUDIT_VERSION,
        "source": {
            "name": source_path.name,
            "sha256": source_hash,
        },
        "sheet": "gwt",
        "raw_rows": len(rows),
        "unique_observations": len(records),
        "valid": flags[QualityFlag.VALID.value],
        "missing": flags[QualityFlag.MISSING.value],
        "duplicate_conflicts": flags[QualityFlag.DUPLICATE.value],
        "station_count": len(station_reports),
        "stations": station_reports,
        "duplicates": duplicate_groups,
        "archived_wrc_rows": wrc_rows,
    }
    return GroundwaterAudit(report=report, records=tuple(records))


def write_groundwater_outputs(
    audit: GroundwaterAudit,
    audit_json_path: str | Path,
    normalized_csv_path: str | Path,
) -> None:
    json_path = Path(audit_json_path)
    csv_path = Path(normalized_csv_path)
    json_path.parent.mkdir(parents=True, exist_ok=True)
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    json_path.write_text(
        json.dumps(audit.report, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    with csv_path.open("w", newline="", encoding="utf-8") as target:
        writer = csv.DictWriter(
            target,
            fieldnames=(
                "observed_date",
                "gwt_m",
                "station_code",
                "region",
                "quality_flag",
                "source_row_number",
            ),
        )
        writer.writeheader()
        for record in audit.records:
            row = asdict(record)
            row["observed_date"] = record.observed_date.isoformat()
            row["quality_flag"] = record.quality_flag.value
            writer.writerow(row)


def import_groundwater(session: Session, path: str | Path) -> dict[str, Any]:
    source_path = Path(path).resolve()
    audit = audit_groundwater(source_path)
    source_hash = audit.report["source"]["sha256"]
    source = session.scalar(select(DataSource).where(DataSource.sha256 == source_hash))
    already_imported = source is not None and session.scalar(
        select(DataImportRun.id).where(
            DataImportRun.source_id == source.id,
            DataImportRun.dataset_kind == "GROUNDWATER",
            DataImportRun.status == "COMPLETED",
        )
    ) is not None
    if source is None:
        source = DataSource(
            name=source_path.name,
            url=source_path.as_posix(),
            sha256=source_hash,
        )
        session.add(source)
        session.flush()

    import_run = DataImportRun(
        source_id=source.id,
        dataset_kind="GROUNDWATER",
        status="STARTED",
        rows_seen=audit.report["unique_observations"],
        rows_inserted=0,
        rows_skipped=0,
        audit_json=audit.report,
    )
    session.add(import_run)
    session.flush()

    if already_imported:
        import_run.status = "COMPLETED"
        import_run.rows_skipped = len(audit.records)
        import_run.completed_at = datetime.now(UTC)
        session.flush()
        return {
            "source_sha256": source_hash,
            "already_imported": True,
            "rows_inserted": 0,
            "rows_skipped": len(audit.records),
        }

    station_rows = [
        {
            "code": station["station_code"],
            "region": station["region"],
            "latitude": None,
            "longitude": None,
            "coordinate_source": None,
        }
        for station in audit.report["stations"]
    ]
    session.execute(
        insert(Station).values(station_rows).on_conflict_do_nothing(index_elements=[Station.code])
    )
    station_ids = dict(session.execute(select(Station.code, Station.id)).all())
    stored_regions = dict(session.execute(select(Station.code, Station.region)).all())
    for station in audit.report["stations"]:
        if stored_regions[station["station_code"]] != station["region"]:
            raise GroundwaterDataError(f"Conflicting region for station {station['station_code']}")
    existing_rows = session.execute(
        select(GroundwaterObservation).where(
            GroundwaterObservation.station_id.in_(station_ids.values())
        )
    ).scalars()
    existing = {(row.station_id, row.observed_date): row for row in existing_rows}
    for record in audit.records:
        previous = existing.get((station_ids[record.station_code], record.observed_date))
        if previous is not None and (
            previous.gwt_m != record.gwt_m
            or previous.quality_flag != record.quality_flag.value
            or previous.mode != MonitoringMode.HISTORICAL_REPLAY.value
        ):
            raise GroundwaterDataError(
                f"Existing observation conflicts with source: "
                f"{record.station_code} {record.observed_date}"
            )
    observation_rows = [
        {
            "station_id": station_ids[record.station_code],
            "observed_date": record.observed_date,
            "gwt_m": record.gwt_m,
            "quality_flag": record.quality_flag.value,
            "mode": MonitoringMode.HISTORICAL_REPLAY.value,
            "source_row_number": record.source_row_number,
            "source_id": source.id,
        }
        for record in audit.records
    ]
    inserted = 0
    for start in range(0, len(observation_rows), 500):
        statement = (
            insert(GroundwaterObservation)
            .values(observation_rows[start : start + 500])
            .on_conflict_do_nothing(
                index_elements=[
                    GroundwaterObservation.station_id,
                    GroundwaterObservation.observed_date,
                ]
            )
            .returning(GroundwaterObservation.id)
        )
        result = session.execute(statement)
        inserted += len(result.all())

    import_run.status = "COMPLETED"
    import_run.rows_inserted = inserted
    import_run.rows_skipped = len(observation_rows) - inserted
    import_run.completed_at = datetime.now(UTC)
    session.flush()
    return {
        "source_sha256": source_hash,
        "already_imported": False,
        "rows_inserted": inserted,
        "rows_skipped": len(observation_rows) - inserted,
    }
